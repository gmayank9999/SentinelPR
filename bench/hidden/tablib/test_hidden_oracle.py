"""Hidden oracle tests for tablib.

Never shown to SentinelPR. The benchmark runs them, with the base revision's own tests,
against every generated PR to label it by execution. They pin behaviour the public suite
leaves loose — separator offsets, duplicate removal, column operations, ragged imports,
format detection, the reST and LaTeX layouts — so a PR can pass every visible test and
still be caught here. Expected values are the behaviour of the released versions used as
bases.
"""

import datetime as dt
from io import BytesIO

import pytest

import tablib
from tablib.core import Row
from tablib.exceptions import HeadersNeeded, InvalidDatasetIndex, InvalidDimensions


def people(headers=True):
    data = tablib.Dataset(headers=["name", "age"] if headers else None)
    data.append(("ada", 36))
    data.append(("alan", 41))
    data.append(("grace", 85))
    return data


class TestDataset:
    def test_repr(self):
        assert repr(tablib.Dataset(title="People")) == "<people dataset>"
        assert repr(tablib.Dataset()) == "<dataset object>"
        assert repr(tablib.Databook()) == "<databook object>"

    def test_row_tags(self):
        row = Row(["x"], tags=["a", "b"])
        assert row.has_tag("a") is True
        assert row.has_tag(None) is False
        assert row.has_tag(["c", "b"]) is True
        assert row.has_tag(["c"]) is False

    def test_filter_by_tag(self):
        data = tablib.Dataset(headers=["n"])
        data.append([1], tags=["odd"])
        data.append([2], tags=["even"])
        data.append([3], tags=["odd"])
        assert [r[0] for r in data.filter("odd")] == [1, 3]
        assert data.height == 3

    def test_push_and_pop_ends(self):
        data = people()
        data.lpush(("zed", 1))
        assert data[0] == ("zed", 1)
        assert data.lpop() == ("zed", 1)
        assert data.rpop() == ("grace", 85)
        assert data.height == 2 and data[0] == ("ada", 36)

    def test_column_operations(self):
        data = people()
        data.lpush_col(["x", "y", "z"], header="id")
        assert data.headers == ["id", "name", "age"]
        assert data[0] == ("x", "ada", 36)
        data.append_col(lambda row: row[2] * 2, header="double")
        assert data["double"] == [72, 82, 170]

    def test_insert_col_into_empty_dataset(self):
        data = tablib.Dataset()
        data.append_col([1, 2, 3])
        assert data.height == 3 and data.width == 1
        assert [r for r in data] == [(1,), (2,), (3,)]

    def test_invalid_dimensions(self):
        data = people()
        with pytest.raises(InvalidDimensions):
            data.append(("only-one",))
        with pytest.raises(InvalidDimensions):
            data.append_col([1, 2])

    def test_separators_offset_by_headers(self):
        data = people()
        data.append_separator("after three")
        assert data._separators[-1] == (4, "after three")
        bare = people(headers=False)
        bare.append_separator("x")
        assert bare._separators[-1] == (3, "x")
        assert tablib.Dataset(headers=["a"])._separators == []
        empty = tablib.Dataset(headers=["a"])
        empty.append_separator()
        assert empty._separators[-1] == (1, "-")
        nothing = tablib.Dataset()
        nothing.append_separator()
        assert nothing._separators[-1] == (0, "-")

    def test_formatter_bounds(self):
        data = people()
        assert data.add_formatter(1, lambda v: v + 1) is True
        assert data.dict[0]["age"] == 37
        with pytest.raises(InvalidDatasetIndex):
            people().add_formatter(3, str)

    def test_sort(self):
        data = people()
        assert [r[0] for r in data.sort("age", reverse=True)] == ["grace", "alan", "ada"]
        assert [r[0] for r in data.sort(0)] == ["ada", "alan", "grace"]
        bare = people(headers=False)
        assert [r[1] for r in bare.sort(1, reverse=True)] == [85, 41, 36]
        with pytest.raises(HeadersNeeded):
            bare.sort("age")

    def test_stack_cols(self):
        left, right = people(), tablib.Dataset(headers=["city"])
        for city in ("london", "paris", "nyc"):
            right.append([city])
        stacked = left.stack_cols(right)
        assert stacked.headers == ["name", "age", "city"]
        assert stacked[2] == ("grace", 85, "nyc")
        with pytest.raises(HeadersNeeded):
            people(headers=False).stack_cols(right)

    def test_remove_duplicates_keeps_first_and_order(self):
        data = tablib.Dataset(headers=["v"])
        for v in (3, 1, 3, 2, 1):
            data.append([v])
        data.remove_duplicates()
        assert data["v"] == [3, 1, 2]

    def test_transpose_without_headers(self):
        data = tablib.Dataset()
        data.append([1, 2])
        data.append([3, 4])
        assert [list(r) for r in data.transpose()] == [[1, 3], [2, 4]]

    def test_databook_size(self):
        book = tablib.Databook([people(), people()])
        assert book.size == 2


class TestFormats:
    def test_csv_ragged_rows_are_padded(self):
        data = tablib.Dataset().load("a,b,c\n1,2\n3,4,5\n", format="csv")
        assert data.headers == ["a", "b", "c"]
        assert data[0] == ("1", "2", "")
        assert data[1] == ("3", "4", "5")

    def test_csv_headerless_first_row_is_data(self):
        data = tablib.Dataset().load("1,2\n3,4\n", format="csv", headers=False)
        assert data.height == 2 and data[0] == ("1", "2")

    def test_csv_export_with_separator_free_data(self):
        assert people().export("csv") == "name,age\r\nada,36\r\nalan,41\r\ngrace,85\r\n"

    @pytest.mark.parametrize(
        ("text", "expected"),
        [("a,b\n1,2\n", "csv"), ('[{"a": 1}]', "json"), ("- {a: 1}\n", "yaml"), ("a\tb\n1\t2\n", "tsv")],
    )
    def test_detect_format(self, text, expected):
        assert tablib.detect_format(text) == expected

    def test_detect_rejects_garbage(self):
        from tablib.formats import registry

        for key in ("json", "xlsx", "xls", "ods", "dbf", "yaml", "csv"):
            assert not registry.get_format(key).detect(BytesIO(b"definitely not tabular {"))

    def test_json_roundtrip_dates(self):
        data = tablib.Dataset(headers=["when"])
        data.append([dt.date(2020, 1, 2)])
        assert data.export("json") == '[{"when": "2020-01-02"}]'
        book = tablib.Databook([people()])
        assert book.export("json").startswith('[{"title": null, "data": [{"name": "ada"')

    def test_yaml_book_export(self):
        book = tablib.Databook([people()])
        assert "name: ada" in book.export("yaml")

    def test_latex_header_rules(self):
        text = people().export("latex")
        assert "\\cmidrule(r){1-1} \\cmidrule(l){2-2}" in text
        assert "\\midrule" not in text
        assert "\\begin{tabular}{lr}" in text
        one = tablib.Dataset(headers=["only"])
        one.append([1])
        assert "    \\midrule\n" in one.export("latex")

    def test_rst_simple_and_grid(self):
        simple = people().export("rst")
        assert simple == "=====  ===\n name  age\n=====  ===\nada    36 \nalan   41 \ngrace  85 \n=====  ==="
        wide = tablib.Dataset(headers=["text"])
        wide.append(["word " * 40])
        grid = wide.export("rst", max_table_width=40)
        assert grid.startswith("+") and all(len(line) <= 44 for line in grid.splitlines())
        book = tablib.Databook([people(), people()])
        assert book.export("rst").count("+") > 0

    def test_xlsx_roundtrip_with_ragged_and_multiline(self):
        data = people()
        data.append(("line\nbreak", 1))
        data.append_separator("end")
        back = tablib.Dataset().load(BytesIO(data.export("xlsx")), format="xlsx")
        assert back.headers == ["name", "age"]
        assert back[0] == ("ada", 36)
        assert back[3] == ("line\nbreak", 1)

    def test_xlsx_import_pads_short_rows(self):
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.append(["a", "b", "c"])
        ws.append([1, 2])
        stream = BytesIO()
        wb.save(stream)
        data = tablib.Dataset().load(stream.getvalue(), format="xlsx")
        assert data[0] == (1, 2, None)

    def test_xlsx_freeze_panes(self):
        from openpyxl import load_workbook

        sheet = load_workbook(BytesIO(people().export("xlsx"))).active
        assert sheet.freeze_panes == "A2"

    def test_ods_roundtrip(self):
        data = people()
        back = tablib.Dataset().load(data.export("ods"), format="ods")
        assert back.headers == ["name", "age"]
        assert back[1] == ("alan", 41)
        book = tablib.Databook().load(tablib.Databook([people()]).export("ods"), format="ods")
        assert book.size == 1

    def test_xls_roundtrip_and_book_headers(self):
        back = tablib.Dataset().load(people().export("xls"), format="xls")
        assert back.headers == ["name", "age"]
        assert back[2] == ("grace", 85.0)
        book = tablib.Databook().load(tablib.Databook([people(), people()]).export("xls"), format="xls")
        assert book.size == 2
        assert book.sheets()[1].headers == ["name", "age"]
